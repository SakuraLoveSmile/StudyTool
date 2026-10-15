from docx import Document
from docx.oxml import OxmlElement
from lxml import etree

from app.services.parser.docx_parser import DocxParser


def test_plain_text_book_contents_headers_appendix_and_sdt_are_preserved(tmp_path):
    book = Document()
    for text in ["书名", "目 录", "第一章 会计基本理论", "第二章 会计报告要素",
                 "第 一 章 会计基本理论", "【学习目标】理解会计。", "第一节 会计信息",
                 "第一，分析审核企业已经发生的经济业务。", "第二，按照权责发生制确认。"]:
        book.add_paragraph(text)
    sdt = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "借：应付账款 2200"
    run.append(text)
    paragraph.append(run)
    content.append(paragraph)
    sdt.append(content)
    book.element.body.insert(len(book.element.body) - 1, sdt)
    table = book.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "资产合计"
    table.cell(0, 1).text = "104000"
    book.add_paragraph("第一章 会计基本理论")  # A repeated running header.
    book.add_paragraph("第一节的后续内容。")
    book.add_paragraph("第 二 章 会计报告要素")
    book.add_paragraph("【学习目标】理解资产。")
    book.add_paragraph("资产是企业拥有或控制的资源。")
    book.add_paragraph("附录一 会计法")
    book.add_paragraph("第一章 总则")
    book.add_paragraph("第一条 为了规范会计行为。")
    path = tmp_path / "book.docx"
    book.save(path)

    chunks = DocxParser.parse_docx(path)
    chapters = list(dict.fromkeys(c["heading"].split(" / ")[0] for c in chunks
                                 if c["heading"].startswith("第")))
    assert chapters == ["第一章 会计基本理论", "第二章 会计报告要素"]
    source = "\n".join(c["content"] for c in chunks)
    assert "第一，分析审核企业已经发生的经济业务。" in source
    assert "第二，按照权责发生制确认。" in source
    assert "借：应付账款 2200" in source
    assert "资产合计 | 104000" in source
    assert all(c["page_number"] is None and c["paragraph_ref"] for c in chunks)
    assert any(c["heading"].startswith("附录一") and "第一条" in c["content"] for c in chunks)
    assert all(not c["heading"].startswith("第一，") for c in chunks)


def test_text_box_text_is_available_without_inventing_page_numbers(tmp_path):
    book = Document()
    book.add_paragraph("第一章 会计基本理论")
    host = book.add_paragraph("图示说明")
    box = OxmlElement("w:txbxContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "会计信息支持管理决策"
    run.append(text)
    paragraph.append(run)
    box.append(paragraph)
    host._p.append(box)
    path = tmp_path / "box.docx"
    book.save(path)
    chunks = DocxParser.parse_docx(path)
    assert sum(c["content"].count("会计信息支持管理决策") for c in chunks) == 1
    assert chunks[0]["heading"] == "第一章 会计基本理论"
    assert all(c["page_number"] is None for c in chunks)


def test_contents_with_page_numbers_does_not_swallow_plain_body(tmp_path):
    book = Document()
    for text in ["目录", "第一章 会计基础 …… 1", "第二章 财务报告 …… 20",
                 "第一章 会计基础", "正文会计信息", "第二章 财务报告", "正文财务报告"]:
        book.add_paragraph(text)
    path = tmp_path / "contents.docx"
    book.save(path)
    chunks = DocxParser.parse_docx(path)
    assert [c["heading"] for c in chunks] == ["目录", "第一章 会计基础", "第二章 财务报告"]


def test_alternate_shape_representations_are_not_both_read(tmp_path):
    book = Document()
    host = book.add_paragraph("图表")
    mc = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
    alternate = etree.Element(mc + "AlternateContent")
    for name in ("Choice", "Fallback"):
        branch = etree.Element(mc + name)
        box = OxmlElement("w:txbxContent")
        paragraph = OxmlElement("w:p")
        run = OxmlElement("w:r")
        text = OxmlElement("w:t")
        text.text = "相同图示的会计信息"
        run.append(text)
        paragraph.append(run)
        box.append(paragraph)
        branch.append(box)
        alternate.append(branch)
    host._p.append(alternate)
    path = tmp_path / "alternate.docx"
    book.save(path)
    chunks = DocxParser.parse_docx(path)
    assert sum(c["content"].count("相同图示的会计信息") for c in chunks) == 1
