#!/usr/bin/env python3
"""Exercise exported images in an isolated network with temporary synthetic data."""
from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import tempfile
import time
import uuid
from pathlib import Path


EMBEDDING_FIXTURE = r'''
import json
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
class Handler(BaseHTTPRequestHandler):
    def do_POST(self):
        body = json.loads(self.rfile.read(int(self.headers['Content-Length'])))
        if self.path != '/v1/embeddings':
            self.send_error(404); return
        inputs = body['input']
        if isinstance(inputs, str): inputs = [inputs]
        payload = json.dumps({'data': [{'index': i, 'embedding': [1.0, 0.5, 0.25, 0.125]}
                                      for i, _ in enumerate(inputs)]}).encode()
        self.send_response(200); self.send_header('Content-Type', 'application/json')
        self.send_header('Content-Length', str(len(payload))); self.end_headers(); self.wfile.write(payload)
    def log_message(self, *args): pass
ThreadingHTTPServer(('0.0.0.0', 9000), Handler).serve_forever()
'''


PROBE = r'''
import io, json, time, urllib.request, urllib.error, http.cookiejar, ssl
from docx import Document
from app.auth import provision_user
origin = 'https://studytool-smoke:8443'
tls = ssl.create_default_context(cafile='/tmp/smoke-ca.pem')
for attempt in range(60):
    try:
        assert urllib.request.urlopen(origin + '/healthz', timeout=2, context=tls).status == 200
        break
    except Exception:
        time.sleep(1)
else:
    raise AssertionError('frontend did not become ready')
assert b'<html' in urllib.request.urlopen(origin + '/library', context=tls).read().lower()
for attempt in range(60):
    try:
        urllib.request.urlopen(origin + '/api/auth/me', timeout=2, context=tls)
    except urllib.error.HTTPError as exc:
        if exc.code == 401:
            break
    except Exception:
        pass
    time.sleep(1)
else:
    raise AssertionError('backend did not become ready through proxy')
provision_user('smoke-alice', 'isolated-start-1234')
provision_user('smoke-bob', 'isolated-start-1234')

def session(username):
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()),
                                        urllib.request.HTTPSHandler(context=tls))
    def call(path, payload=None, csrf=None, expected=200, raw=None, content_type=None):
        headers = {'Origin': origin}
        data = raw
        if payload is not None:
            data = json.dumps(payload).encode()
            headers['Content-Type'] = 'application/json'
        if content_type:
            headers['Content-Type'] = content_type
        if csrf:
            headers['X-CSRF-Token'] = csrf
        request = urllib.request.Request(origin + path, data=data, headers=headers)
        try:
            response = opener.open(request, timeout=10)
        except urllib.error.HTTPError as exc:
            response = exc
        assert response.status == expected, (path, response.status, response.read().decode())
        body = response.read()
        return json.loads(body) if body else None
    login = call('/api/auth/login', {'username': username, 'password': 'isolated-start-1234'})
    changed = call('/api/auth/change-password', {'current_password': 'isolated-start-1234',
                   'new_password': 'isolated-changed-1234'}, login['csrf_token'])
    return call, changed['csrf_token']

try:
    urllib.request.urlopen(origin + '/api/library/guides', context=tls)
    raise AssertionError('anonymous library access allowed')
except urllib.error.HTTPError as exc:
    assert exc.code == 401
alice, token = session('smoke-alice')
bob, _ = session('smoke-bob')
assert alice('/api/library/guides') == []
assert alice('/api/library/answers') == []
alice('/api/library/ask', {'question': 'test', 'doc_ids': []}, expected=403)
document = Document()
document.add_heading('第一章 会计基础', level=1)
document.add_paragraph('会计信息用于经济决策，管理活动需要真实完整的会计信息。' * 30)
stream = io.BytesIO(); document.save(stream)
boundary = 'studytool-isolated-smoke'
body = (f'--{boundary}\r\nContent-Disposition: form-data; name="file_type"\r\n\r\ntextbook\r\n'
        f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="smoke.docx"\r\n'
        'Content-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n').encode()
body += stream.getvalue() + f'\r\n--{boundary}--\r\n'.encode()
uploaded = alice('/api/materials/upload', csrf=token, raw=body,
                 content_type=f'multipart/form-data; boundary={boundary}')
doc_id = uploaded['id']
for attempt in range(60):
    docs = alice('/api/materials')
    current = next(d for d in docs if d['id'] == doc_id)
    if current['status'] == 'completed':
        break
    if current['status'] == 'failed':
        raise AssertionError(current)
    time.sleep(1)
else:
    raise AssertionError('DOCX parse did not complete')
chapters = alice('/api/library/chapters?doc_id=' + doc_id)
assert chapters and chapters[0]['character_count'] > 0
assert bob('/api/materials') == []
bob('/api/library/chapters?doc_id=' + doc_id, expected=404)
assert alice('/api/auth/me')['user']['username'] == 'smoke-alice'
print(json.dumps({'result': 'passed', 'checks': ['spa-deep-link', 'anonymous-denied', 'login',
    'password-change', 'csrf', 'docx-upload-parse', 'library-chapters', 'tenant-isolation'],
    'remote_model_called': False, 'embedding_protocol_fixture': True}))
'''


RESTART_PROBE = r'''
import http.cookiejar, json, ssl, time, urllib.request
origin = 'https://studytool-smoke:8443'
tls = ssl.create_default_context(cafile='/tmp/smoke-ca.pem')
for _ in range(60):
    try:
        urllib.request.urlopen('http://127.0.0.1:8000/health', timeout=2)
        break
    except Exception: time.sleep(1)
else: raise AssertionError('backend did not restart')
opener = urllib.request.build_opener(urllib.request.HTTPSHandler(context=tls),
    urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))
request = urllib.request.Request(origin + '/api/auth/login',
    data=json.dumps({'username': 'smoke-alice', 'password': 'isolated-changed-1234'}).encode(),
    headers={'Content-Type': 'application/json', 'Origin': origin})
with opener.open(request, timeout=10) as response:
    assert json.load(response)['user']['username'] == 'smoke-alice'
with opener.open(origin + '/api/materials') as response:
    docs = json.load(response)
assert len(docs) == 1 and docs[0]['status'] == 'completed'
with opener.open(origin + '/api/library/chapters?doc_id=' + docs[0]['id']) as response:
    assert json.load(response)
print(json.dumps({'result': 'passed', 'checks': ['backend-restart', 'account-persistence',
    'material-persistence', 'chapter-persistence'], 'remote_model_called': False}))
'''


def run(*args: str, **kwargs):
    return subprocess.run(args, check=True, text=True, **kwargs)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--bundle', type=Path, required=True)
    args = parser.parse_args()
    metadata = json.loads((args.bundle / 'BUILD.json').read_text())
    suffix = uuid.uuid4().hex[:10]
    network, backend, frontend, model = (f'studytool-smoke-{suffix}-{name}' for name in ('net', 'backend', 'frontend', 'fixture'))
    created = []
    volume = f'studytool-smoke-{suffix}-data'
    fixture = tempfile.TemporaryDirectory(prefix='studytool-container-smoke-')
    fixture_config = Path(fixture.name) / 'Caddyfile'
    fixture_config.write_text((Path(__file__).parent / 'Caddyfile').read_text().replace(
        '{$DOMAIN} {', '{$DOMAIN} {\n    tls internal'))
    run('docker', 'network', 'create', '--internal', network)
    run('docker', 'volume', 'create', '--label', 'studytool.disposable-smoke=true', volume)
    try:
        run('docker', 'run', '-d', '--platform', metadata['platform'], '--name', model, '--network', network,
            '--network-alias', 'model-fixture', '--entrypoint', 'python',
            metadata['images']['backend'], '-u', '-c', EMBEDDING_FIXTURE)
        created.append(model)
        run('docker', 'run', '-d', '--platform', metadata['platform'], '--name', backend, '--network', network, '--network-alias', 'backend',
            '-v', f'{volume}:/data',
            '-e', 'PUBLIC_MODE=true', '-e', 'PUBLIC_ORIGIN=https://studytool-smoke:8443',
            '-e', 'AUTH_COOKIE_SECURE=true', '-e', 'OCR_MODE=vision_llm',
            '-e', 'DATA_DIR=/data', '-e', 'DEMO_MODE=false',
            '-e', 'EMBEDDING_MODE=remote', '-e', 'EMBEDDING_BASE_URL=http://model-fixture:9000/v1',
            '-e', 'EMBEDDING_API_KEY=isolated-fixture',
            '-e', 'LLM_BASE_URL=http://127.0.0.1:9/v1', '-e', 'LLM_API_KEY=smoke-unused',
            metadata['images']['backend'])
        created.append(backend)
        run('docker', 'run', '-d', '--platform', metadata['platform'], '--name', frontend, '--network', network,
            '--network-alias', 'studytool-smoke', '-e', 'DOMAIN=https://studytool-smoke:8443',
            '-v', f'{fixture_config}:/etc/caddy/Caddyfile:ro',
            metadata['images']['frontend'])
        created.append(frontend)
        # Trust the fixture CA only inside this disposable Python probe. Never
        # bypass TLS or install a trusted root on the host/user's devices.
        for _ in range(30):
            ca = subprocess.run(['docker', 'exec', frontend, 'cat',
                '/data/caddy/pki/authorities/local/root.crt'], capture_output=True, text=True)
            if ca.returncode == 0:
                break
            time.sleep(1)
        else:
            raise RuntimeError('isolated Caddy CA was not created')
        run('docker', 'exec', '-i', backend, 'python', '-c',
            'import sys; open("/tmp/smoke-ca.pem", "w").write(sys.stdin.read())',
            input=ca.stdout, capture_output=True)
        result = run('docker', 'exec', '-i', backend, 'python', '-', input=PROBE, capture_output=True)
        print(result.stdout)
        run('docker', 'restart', backend, capture_output=True)
        restarted = run('docker', 'exec', '-i', backend, 'python', '-', input=RESTART_PROBE, capture_output=True)
        print(restarted.stdout)
        evidence = {'initial': json.loads(result.stdout.strip().splitlines()[-1]),
                    'restart': json.loads(restarted.stdout.strip().splitlines()[-1]),
                    'platform': metadata['platform']}
        evidence_text = json.dumps(evidence, indent=2) + '\n'
        (args.bundle / 'SMOKE.json').write_text(evidence_text)
        checksums = args.bundle / 'SHA256SUMS.txt'
        entries = [line for line in checksums.read_text().splitlines() if not line.endswith('  SMOKE.json')]
        entries.append(hashlib.sha256(evidence_text.encode()).hexdigest() + '  SMOKE.json')
        checksums.write_text('\n'.join(entries) + '\n')
    except subprocess.CalledProcessError as exc:
        if exc.stdout: print(exc.stdout)
        if exc.stderr: print(exc.stderr)
        for container in created:
            subprocess.run(['docker', 'logs', '--tail', '40', container], check=False)
        raise
    finally:
        for container in reversed(created):
            subprocess.run(['docker', 'rm', '-f', container], check=False, capture_output=True)
        subprocess.run(['docker', 'network', 'rm', network], check=False, capture_output=True)
        subprocess.run(['docker', 'volume', 'rm', volume], check=False, capture_output=True)
        fixture.cleanup()


if __name__ == '__main__':
    main()
